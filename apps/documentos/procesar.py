"""Lectura automática de datos de los archivos subidos: fotos, planillas y nombres de planos."""

import csv
import hashlib
import io
import logging
import re
from datetime import datetime

from django.core.files.base import ContentFile
from django.utils import timezone

log = logging.getLogger(__name__)

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:  # sin soporte HEIC: las fotos de iPhone solo se descargan
    pass

LADO_VISTA_PREVIA = 1600
FILAS_VISTA_PREVIA = 50
COLUMNAS_VISTA_PREVIA = 20
HOJAS_VISTA_PREVIA = 5


def huella(archivo):
    """SHA-256 de un archivo subido, leyéndolo por partes."""
    h = hashlib.sha256()
    for parte in archivo.chunks():
        h.update(parte)
    archivo.seek(0)
    return h.hexdigest()


# --- Nombres de planos ------------------------------------------------------

_PLANO = re.compile(
    r"^(?P<numero>[A-Za-z]{1,3}[-_ ]?\d{1,4}[A-Za-z]?)"
    r"(?:[\s_\-]+(?P<titulo>.*?))?"
    r"(?:[\s_\-]*(?:rev(?:isi[oó]n)?)\.?[\s_\-]*(?P<rev>[A-Za-z]{1,2}|\d{1,2}))?$",
    re.IGNORECASE,
)


def leer_nombre_plano(nombre_archivo):
    """'A-101_Planta baja_RevC.pdf' -> {'numero': 'A-101', 'titulo': 'Planta baja', 'revision': 'C'}."""
    base = re.sub(r"\.[^.]+$", "", nombre_archivo).strip()
    m = _PLANO.match(base)
    if not m:
        return {}
    numero = re.sub(r"[_ ]", "-", m.group("numero").upper())
    if "-" not in numero:
        numero = re.sub(r"^([A-Z]+)(\d)", r"\1-\2", numero)
    titulo = (m.group("titulo") or "").replace("_", " ").strip(" -")
    return {"numero": numero, "titulo": titulo, "revision": (m.group("rev") or "").upper()}


def titulo_desde_nombre(nombre_archivo):
    base = re.sub(r"\.[^.]+$", "", nombre_archivo)
    return re.sub(r"[_]+", " ", base).strip()[:200] or nombre_archivo[:200]


# --- Fotos -------------------------------------------------------------------

def procesar_imagen(version):
    from PIL import Image, ImageOps

    with version.archivo.open("rb") as f:
        imagen = Image.open(f)
        imagen.load()
    exif = imagen.getexif()
    fecha = exif.get_ifd(0x8769).get(36867) or exif.get(306)  # DateTimeOriginal o DateTime
    if fecha:
        try:
            version.fecha_toma = timezone.make_aware(datetime.strptime(str(fecha), "%Y:%m:%d %H:%M:%S"))
        except ValueError:
            pass
    imagen = ImageOps.exif_transpose(imagen).convert("RGB")
    imagen.thumbnail((LADO_VISTA_PREVIA, LADO_VISTA_PREVIA))
    salida = io.BytesIO()
    imagen.save(salida, "JPEG", quality=82, optimize=True)
    version.vista_previa.save("vista.jpg", ContentFile(salida.getvalue()), save=False)
    version.info = {"ancho": imagen.width, "alto": imagen.height}


# --- Planillas --------------------------------------------------------------

def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, int):
        return f"{valor:,}".replace(",", ".")
    if isinstance(valor, float):
        texto = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return texto[:-3] if texto.endswith(",00") else texto
    return str(valor)


def _recortar(datos):
    """Saca las filas y columnas vacías del final."""
    while datos and not any(datos[-1]):
        datos.pop()
    ancho = max((max((i + 1 for i, c in enumerate(fila) if c), default=0) for fila in datos), default=0)
    return [fila[:ancho] + [""] * (ancho - len(fila[:ancho])) for fila in datos]


def leer_planilla(version, filas=FILAS_VISTA_PREVIA):
    """[{'nombre': hoja, 'filas': total, 'columnas': total, 'datos': [[...], ...]}] con las primeras filas."""
    if version.extension == "csv":
        with version.archivo.open("rb") as f:
            texto = f.read().decode("utf-8-sig", errors="replace")
        separador = ";" if texto[:2000].count(";") > texto[:2000].count(",") else ","
        todas = list(csv.reader(io.StringIO(texto), delimiter=separador))
        datos = _recortar([fila[:COLUMNAS_VISTA_PREVIA] for fila in todas[:filas]])
        return [{"nombre": "CSV", "filas": len(todas), "columnas": max((len(f) for f in todas), default=0),
                 "datos": datos}]

    from openpyxl import load_workbook

    with version.archivo.open("rb") as f:
        libro = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    hojas = []
    for hoja in libro.worksheets[:HOJAS_VISTA_PREVIA]:
        datos = []
        for fila in hoja.iter_rows(max_row=filas, max_col=COLUMNAS_VISTA_PREVIA, values_only=True):
            datos.append([_texto(v) for v in fila])
        hojas.append({"nombre": hoja.title, "filas": hoja.max_row or 0, "columnas": hoja.max_column or 0,
                      "datos": _recortar(datos)})
    libro.close()
    return hojas


def procesar_version(version):
    """Lee los datos que se puedan del archivo. Nunca frena la subida: si algo falla, se registra y sigue."""
    try:
        if version.es_imagen:
            procesar_imagen(version)
        elif version.es_planilla:
            hojas = leer_planilla(version, filas=0)
            version.info = {"hojas": [{"nombre": h["nombre"], "filas": h["filas"]} for h in hojas]}
    except Exception:  # noqa: BLE001 - un archivo dañado no tiene que romper la carga
        log.warning("No se pudieron leer los datos de %s", version.nombre_original, exc_info=True)
        version.info = {**version.info, "aviso": "No se pudo leer el contenido del archivo."}
